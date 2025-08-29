import re
from collections.abc import MutableSequence, Set, Sequence
from enum import IntEnum, Enum
from typing import NamedTuple

from docx import Document
from docx.table import Table

from sdp_lib.passport.mixins import ReprMixin


class TableCategories(IntEnum):
    directions = 0
    time_program_va = 1
    time_program_ft = 2

    @classmethod
    def get_tp_vals(cls) -> Set:
        return {cls.time_program_ft, cls.time_program_va}


class TableMeta(NamedTuple):
    index: int
    category: TableCategories


class DocTablesMeta(ReprMixin):
    def __init__(self, tables: Sequence[TableMeta]):
        self._tables = tables
        self._directions = None
        self._ft = []
        self._va = []
        self.sort_tables_by_categories()

    def __iter__(self):
        return (k_v for k_v in self._tables)

    def __getitem__(self, item):
        return self._tables[item]

    def __len__(self):
        return len(self._tables)

    @property
    def directions(self) -> int:
        return self._directions

    @property
    def va_tables(self) -> MutableSequence[int]:
        return self._va

    @property
    def num_tables(self) -> int:
        return len(self._tables)

    def sort_tables_by_categories(self):
        for table in (table for table in self._tables if table.category is not None):
            if table.category == TableCategories.directions:
                self._directions = table.index
            elif table.category == TableCategories.time_program_ft:
                self._ft.append(table.index)
            elif table.category == TableCategories.time_program_va:
                self._va.append(table.index)

    def get_tables(self):
        return self._tables


class DocTable:
    def __init__(self, index, head, rows):
        self._index = index
        self._head = head
        self._rows = rows

"""
Таблица направлений.
length = 14
row0 = ['№ нап.', 'Тип направления', 'Фазы, в кот. участ. направ.', 'Светофоры', '"Запрет"', '"Запрет"', '"Запрет"', '"Запрет"', '"Разрешение"', '"Разрешение"', 'Пост. красное', 'ТООВ ', 'ТООВ ', 'Примечание']
row1 = ['№ нап.', 'Тип направления', 'Фазы, в кот. участ. направ.', 'Светофоры', 'Тзд', 'Тзм', 'Тж', 'Тк', 'Ткж', 'Тз', 'Пост. красное', 'Красн.', 'Зелен.', '']
---------------
Временная программа фикс.
length = 10
row0 = ['Программа 1', 'Программа 1', 'Программа 1', 'Тц = 89 с.', 'пн,вт,ср,чт,пт,сб,вс:', 'пн,вт,ср,чт,пт,сб,вс:', '06:00:00-21:00:00', '06:00:00-21:00:00', '06:00:00-21:00:00', '06:00:00-21:00:00']
row1 = ['№ пп', '№ фазы', 'Направления', 'Направления', 'Направления', 'Тосн', 'Тосн', 'Тсдвиг', 'Тмин', 'Тип фазы']

Временная программа адаптива.
length = 11
row0 = ['Программа 1', 'Программа 1', '', '', 'пн,вт,ср,чт,пт,сб,вс:', 'пн,вт,ср,чт,пт,сб,вс:', '', '', '06:00:00-21:00:00', '06:00:00-21:00:00', '06:00:00-21:00:00']
row1 = ['№ фазы', 'Направления', 'Направления', 'Условия вызова', 'Условия продления', 'Условия завершения', 'Условия завершения', 'Тмин', 'Тмин', 'Тмакс1\nТр.', 'Тмакс2\nОТ']

"""


class DirectionTablePatterns(Enum):
    row0_cell0 = re.compile('^№\s*нап', re.IGNORECASE)
    row0_cell1 = re.compile('^тип\s*направления', re.IGNORECASE)


class TimeProgramFtPatterns(Enum):
    row1_cell0 = re.compile('№\s*пп',  re.IGNORECASE)
    row1_cell1 = re.compile('№\s*фазы',  re.IGNORECASE)
    row1_cell9 = re.compile('Тип\s*фазы',  re.IGNORECASE)


class TimeProgramVaPatterns(Enum):
    row1_cell0 = re.compile('№\s*фазы',  re.IGNORECASE)
    row1_cell1 = re.compile('направления',  re.IGNORECASE)
    row1_cell10 = re.compile('макс.*2',  re.IGNORECASE)


def _check_is_directions_table(rows) -> bool:
    return bool(
        14 <= len(rows[0].cells) <= 15
        and re.search(DirectionTablePatterns.row0_cell0.value, rows[0].cells[0].text) is not None
        and re.search(DirectionTablePatterns.row0_cell1.value, rows[0].cells[1].text) is not None
    )


def _check_is_time_program_table_ft(rows) -> bool:
    return bool(
        len(rows[0].cells) == 10
        and re.search(TimeProgramFtPatterns.row1_cell0.value, rows[1].cells[0].text) is not None
        and re.search(TimeProgramFtPatterns.row1_cell1.value, rows[1].cells[1].text) is not None
        and re.search(TimeProgramFtPatterns.row1_cell9.value, rows[1].cells[9].text) is not None
    )


def _check_is_time_program_table_va(rows) -> bool:
    return bool(
        len(rows[0].cells) == 11
        and re.search(TimeProgramVaPatterns.row1_cell0.value, rows[1].cells[0].text) is not None
        and re.search(TimeProgramVaPatterns.row1_cell1.value, rows[1].cells[1].text) is not None
        and re.search(TimeProgramVaPatterns.row1_cell10.value, rows[1].cells[10].text) is not None
    )

    # if len(cells) == 10:
    #     if all(re.findall(pattern, string) for pattern, string in (
    #             (first_cell_time_program_pattern, cells[0]),
    #             (last_cell_time_program_pattern, cells[len(cells) - 1]),
    #     )):
    #         return True
    # return False


def sort(tables: MutableSequence[Table]) -> DocTablesMeta:
    tables_meta = []
    for i, table in enumerate(tables):
        t_rows = table.rows
        entity = None
        if _check_is_directions_table(t_rows):
            entity = TableCategories.directions
        elif _check_is_time_program_table_ft(t_rows):
            entity = TableCategories.time_program_ft
        elif _check_is_time_program_table_va(t_rows):
            entity = TableCategories.time_program_va
        tables_meta.append(TableMeta(i, entity))
    return DocTablesMeta(tables_meta)


def _display_all_tables(doc_x):
    """ Выводит на экран данные всех таблиц doc(x) файла. """
    for table in doc_x.tables:
        print(f'-- Start Table --')
        print(f'table.rows: {list(table.rows)}')
        print(f'table.row: {table.rows}')
        print(f'len(rows): {len(table.rows)}')
        print(f'len(cells): {len(table.rows[0].cells)}')
        for i, row in enumerate(table.rows):
            print(f'{i}: {[cell.text for cell in row.cells]}')
        print(f'*' * 100)


if __name__ == '__main__':
    doc = Document('/home/auser/Downloads/СО 20250006 ул. Островитянова,д, 15.docx')
    # print(doc.tables)
    _display_all_tables(doc)
    print(sort(doc.tables))


