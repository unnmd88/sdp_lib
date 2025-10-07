import itertools
import re
from collections.abc import Iterable
from dataclasses import dataclass, asdict, fields, astuple
from enum import IntEnum
from typing import NamedTuple

from docx import Document

if __name__ == '__main__':
    # doc = Document('Новый текстовый документ.txt')
    # doc = Document('Новый текстовый документ.zip')
    # doc = Document('Новый текстовый документ — копия.docx')
    doc = Document('1.jpg')
    print(doc)


