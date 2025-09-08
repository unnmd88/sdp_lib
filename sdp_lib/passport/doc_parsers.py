import re
from collections.abc import MutableSequence
from enum import IntEnum
from typing import NamedTuple

from docx import Document
from docx.table import Table

from sdp_lib.passport.mixins import ReprMixin


class Tables(IntEnum):
    directions = 0
    time_programs = 1


class TableMeta(NamedTuple):
    index: int
    category: Tables


class DocTables(ReprMixin):
    def __init__(self):
        self._tables = {}
        self._tables_by_categories = {}

    def __iter__(self):
        return (k_v for k_v in self._tables.items())

    def __getitem__(self, item):
        return self._tables[item]

    def __setitem__(self, key, value):
        self._tables[key] = value

    def load_table_data(self, data: TableMeta):
        self._tables[data.index] = data.category

    def get_tables(self):
        return self._tables


class DocTable:
    def __init__(self, index, head, rows):
        self._index = index
        self._head = head
        self._rows = rows



doc = Document('/home/auser/Downloads/СО 20250006 ул. Островитянова,д, 15.docx')
print(doc.tables)
print(doc.tables[0].rows)


directions_dir_tbl_pattern = re.compile('^тип\s*направл', re.IGNORECASE)
tzd_dir_tbl_pattern = re.compile('^тзд', re.IGNORECASE)

first_cell_time_program_pattern = re.compile('№\s*фазы|№\s*пп',  re.IGNORECASE)
last_cell_time_program_pattern = re.compile('тип\s*фазы|макс2',  re.IGNORECASE)



def check_is_directions_table(cells) -> bool:
    return bool(
        (len(cells) == 14
        and re.findall(directions_dir_tbl_pattern, cells[1])
        and re.findall(tzd_dir_tbl_pattern, cells[4]))
    )

def check_is_time_program_table(cells) -> bool:
    if len(cells) >= 10:
        if all(re.findall(pattern, string) for pattern, string in (
                (first_cell_time_program_pattern, cells[0]),
                (last_cell_time_program_pattern, cells[len(cells) - 1]),
        )):
            return True
    return False


def identify_table_entity(table_rows) -> Tables:
    for row in table_rows:
        cells = [cell.text_is_valid for cell in row.cells]
        if check_is_directions_table(cells):
            print(cells)
            return Tables.directions
        elif check_is_time_program_table(cells):
            print(cells)
            return Tables.time_programs
        else:
            pass
    return None


def sort_tables(tables: MutableSequence[Table]) -> DocTables:
    tbls = DocTables()
    for i, table in enumerate(tables):
        tbls.load_table_data(TableMeta(i, identify_table_entity(table.rows)))
    return tbls


# for row in doc.tables[0].rows:
#     cells = [cell.text for cell in row.cells]
#     if re.findall(directions_pattern, cells[1]) and re.findall(tzd_pattern, cells[4]):
#         print(f'len(cells): {len(cells)}')
#         print(cells)
#         for r in row:
#             _cells = [cell.text for cell in row.cells]
#             print(cells)
    # for cell in row.cells:
    #     print(repr(cell.text))


if __name__ == '__main__':
    sorted_tables = sort_tables(doc.tables)
    print(sorted_tables)
    for k, v in sorted_tables:
        if v:
            for row in doc.tables[k].rows:
                # cells = [cell.text for cell in row.cells]
                print('|'.join(cell.text_is_valid for cell in row.cells))