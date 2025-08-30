import itertools
import os.path
import re
from collections import Counter
from collections.abc import Sequence, Iterable
from dataclasses import make_dataclass
from enum import IntEnum
from pathlib import Path
from typing import NamedTuple

from docx.table import _Row

from sdp_lib.passport.constants import DirectionTypes, RowNames, StorageNames, ColNamesDirectionsTable, \
    MessageCategories, TableNames, Fields
from sdp_lib.passport.mixins import ReprMixin
from sdp_lib.passport.text_messages import Text
from sdp_lib.passport.passport2.base import Cell, StageOrDirectionCell, get_cell_with_value_as_number_, get_cell, \
    get_cell_with_value_as_prom_tact_time, Message, AbstractRow, AbstractTableWithStages, InitData
# from sdp_lib.passport.passport2.doc_parsers.tables_sorting import sort
from sdp_lib.utils_common.utils_common import to_json


class _Types(IntEnum):
    all_fields15 = 15
    exclude_tzz_fields14 = 14


def _get_fields(fields_type: _Types):
    all_fields = (
        ('number', Cell),
        ('direction_type', Cell),
        ('stages', StageOrDirectionCell),
        ('traffic_lights', Cell),
        ('t_green_ext', Cell),
        ('t_flashing_green', Cell),
        ('t_yellow', Cell),
        ('t_red', Cell),
        ('t_red_yellow', Cell),
        ('t_z', Cell),
        ('t_zz', Cell),
        ('always_red', Cell),
        ('toov_red', Cell),
        ('toov_green', Cell),
        ('description', Cell)
    )
    if fields_type == _Types.exclude_tzz_fields14:
        return tuple(f for f in all_fields if f[0] != 't_zz')
    return all_fields


DirectionRowCellsAll = make_dataclass(
    'DirectionRowCellsAll',
    _get_fields(_Types.all_fields15),
    slots=True,
    frozen=True
)


class DirectionRowCells(NamedTuple):
    number: Cell
    direction_type: Cell
    stages: StageOrDirectionCell
    traffic_lights: Cell
    t_green_ext: Cell
    t_flashing_green: Cell
    t_yellow: Cell
    t_red: Cell
    t_red_yellow: Cell
    t_z: Cell
    t_zz: Cell
    always_red: Cell
    toov_red: Cell
    toov_green: Cell
    description: Cell


# @dataclass(frozen=True, slots=True)
# class DirectionRowCells:
#     number: Cell
#     direction_type: Cell
#     stages: StageOrDirectionCell
#     traffic_lights: Cell
#     t_green_ext: Cell
#     t_flashing_green: Cell
#     t_yellow: Cell
#     t_red: Cell
#     t_red_yellow: Cell
#     t_z: Cell
#     t_zz: Cell
#     always_red: Cell
#     toov_red: Cell
#     toov_green: Cell
#     description: Cell


class RowStructure(IntEnum):
    number           = 0
    direction_type   = 1
    stages           = 2
    traffic_lights   = 3
    t_green_ext      = 4
    t_flashing_green = 5
    t_yellow         = 6
    t_red            = 7
    t_red_yellow     = 8
    t_z              = 9
    t_zz             = 10
    is_always_red    = 11
    toov_red         = 12
    toov_green       = 13
    description      = 14


class DirectionRow(AbstractRow, ReprMixin):

    name = RowNames.direction
    standard_directions = re.compile('Транспорт|Поворотное|Пешеходное|общ.*тр|пос.*крас', re.IGNORECASE)

    ALWAYS_RED = re.compile(r'кр|-', re.IGNORECASE)

    def __init__(
            self, index: int,
            row: _Row,
            is_empty: bool,
            is_header: bool,
            number: InitData = InitData(),
            direction_type: InitData = InitData(),
            stages: InitData = InitData(),
            traffic_lights: InitData = InitData(),
            t_green_ext: InitData = InitData(),
            t_flashing_green: InitData = InitData(),
            t_yellow: InitData = InitData(),
            t_red: InitData = InitData(),
            t_red_yellow: InitData = InitData(),
            t_z: InitData = InitData(),
            t_zz: InitData = InitData(),
            always_red: InitData = InitData(),
            toov_red: InitData = InitData(),
            toov_green: InitData = InitData(),
            description:InitData = InitData(),

    ):
        super().__init__(index, row, is_empty, is_header)

        _stages = StageOrDirectionCell(*stages)
        dt = get_cell(*direction_type)
        dt_val = dt.value
        self._direction_type_is_standard = True if re.search(self.standard_directions, dt_val) is not None else False
        self._cells = DirectionRowCells(
            get_cell_with_value_as_number_(*number),
            dt,
            _stages,
            get_cell(*traffic_lights),
            get_cell_with_value_as_prom_tact_time(*t_green_ext, dt_val, ColNamesDirectionsTable.t_green_ext),
            get_cell_with_value_as_prom_tact_time(*t_flashing_green, dt_val, ColNamesDirectionsTable.t_flashing_green,),
            get_cell_with_value_as_prom_tact_time(*t_yellow, dt_val, ColNamesDirectionsTable.t_yellow,),
            get_cell_with_value_as_prom_tact_time(*t_red, dt_val, ColNamesDirectionsTable.t_red,),
            get_cell_with_value_as_prom_tact_time(*t_red_yellow, dt_val, ColNamesDirectionsTable.t_red_yellow ),
            get_cell_with_value_as_prom_tact_time(*t_z, dt_val, ColNamesDirectionsTable.t_z),
            get_cell_with_value_as_prom_tact_time(*t_zz, dt_val, ColNamesDirectionsTable.t_zz),
            get_cell(*always_red),
            get_cell(*toov_red),
            get_cell(*toov_green),
            get_cell(*description)
            )

        if not self.cells.number.is_valid:
            self._extra_data.err_and_warn.add_errors(
                Message(Text.get_bad_num(self._cells.number.value, ColNamesDirectionsTable.number), MessageCategories.validation),
            )
            self._extra_data.permissions.set_val_for_compare_stages(False)
        if not self.cells.stages.is_valid:
            self._extra_data.err_and_warn.add_errors(
                Message(Text.get_bad_val(self._cells.stages.value, ColNamesDirectionsTable.stages), MessageCategories.validation)
            )
            self._extra_data.permissions.set_val_for_compare_stages(False)



    def _get_values_to_set_in_cells(self) -> Iterable:
        if len(self._row.cells) == _Types.exclude_tzz_fields14:
            return itertools.chain(
                (self._row.cells[i].text for i in range(10)),
                (None, ),
                (self._row.cells[i].text for i in range(10, _Types.exclude_tzz_fields14))
            )
        return (cell.text for cell in self._cells)

        # for i, cell in enumerate(self._row.cells):
        #     yield cell.text
        #     if i == 10 and len(self._row.cells) == _Types.exclude_tzz_fields14:
        #         yield None

    def _get_direction_type(self, init_val: str | DirectionTypes, pos: int | None) -> Cell:
        default_val, is_valid = DirectionTypes.empty, True
        if re.findall(self.ALWAYS_RED, init_val):
            val = DirectionTypes.always_red
        elif init_val:
            val = init_val
            try:
                DirectionTypes(init_val)
            except ValueError:
                is_valid = False
                self._extra_data.err_and_warn.add_warnings(
                    Message(
                        f'Задан нестандартный тип направления: {init_val}. '
                        f'Стандартные типы: {DirectionTypes.get_standard_types()}',
                        MessageCategories.validation
                    )
                )
        else:
            val = default_val
        return Cell(pos, init_val, default_val, val, is_valid)

    def _check_direction_type_is_standard(self) -> bool:
        try:
            return bool(DirectionTypes(self.cells.direction_type.value))
        except ValueError:
            self._extra_data.err_and_warn.add_warnings(
                Message('Задан нестандартный тип направления.', MessageCategories.validation)
            )
            return False


class DirectionsTable(AbstractTableWithStages, ReprMixin):

    name = TableNames.directions_table
    allowed_cnt_row_props = {1, 3, 14, 15}
    row_class = DirectionRow
    key_name = Fields.directions
    start_vals_row = 2

    # def __init__(self, index, rows: Sequence[_Row], rows2: Sequence[DirectionRow]):
    #     super().__init__(index, rows, rows2)
    #     self._direction_type_counter = Counter(str(row.cells.direction_type.value) for row in self._rows)

    def get_direction_types_cnt(self):
        return Counter(str(row.cells.direction_type.value) for row in self._rows)

    def dump_to_dict(self) -> dict:
        return super().dump_to_dict() | {
            str(Fields.directions_by_type): {k: v for k, v in self.get_direction_types_cnt().items()}
        }

if __name__ == '__main__':
    from docx import Document
    # print(DirectionRowCellsExcludeTzz(*['№ нап.', 'Тип направления', 'Фазы, в кот. участ. направ.', 'Светофоры', '"Запрет"', '"Запрет"', '"Запрет"', '"Запрет"', '"Разрешение"', '"Разрешение"', 'Пост. красное', 'ТООВ ', 'ТООВ ', 'Примечание']))
    doc = Document('C://Programms//py.projects//sdp_lib//sdp_lib//passport//СО_2094_ул_Островитянова_ул_Ак_Волгина (2).docx')
    # sorted_tables = sort(doc.tables)
    # r = DirectionRow(
    #     0, doc.tables[sorted_tables.directions].rows[3],
    #     all(not v.text for v in doc.tables[sorted_tables.directions].rows[3].cells),
    #     False
    # )
    # print(to_json(r.dump_to_dict()))
    # dt = DirectionsTable(sorted_tables.directions, doc.tables[sorted_tables.directions].rows)
    #
    # print(to_json(dt.dump_to_dict()))