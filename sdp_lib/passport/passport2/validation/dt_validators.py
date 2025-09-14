import itertools
import re
from collections.abc import Generator, Sequence
from enum import IntEnum
from typing import Any, NamedTuple

from docx import Document
from docx.shared import Length, Inches, RGBColor
from docx.table import (
    _Rows,
    Table
)

from sdp_lib.passport.constants import (
    allowed_column_lengths_dt,
    allowed_min_num_rows,
    DirectionEntities,
    PatternsDirectionTable,
    HeadRowsDirectionTableData,
    # dt_mapping_from_length,
    dt_timing_columns_mapping,
    timing_matches,
    AllowedValues,
    substring_for_search_tlc, toov_matches, mapping_direction_data, DirectionDataContainer, head_rows_data_dt
)
from sdp_lib.passport.passport2.base2 import (
    DirectionDataRow,
    CellData,
    NumberValidation,
    CellMapping,
    MessageStorage, TheTable
)
from sdp_lib.passport.passport2.utils import (
    remove_left_light_spaces_from_cell_text,
    found_pos_start_num,
    repair_string_if_sep_in_illegal_pos, write_messages_to_cell
)
from sdp_lib.passport.passport2.validation.common_validators import (
    validate_geometry,
    validate_sequence_directions_or_stages_nums_and_create_cell,
    match_one_string_to_many_patterns_and_get_alias_and_create_cell,
    create_cells_for_head_row,
    num_validate_and_create_cell, lrstrip_in_cell_and_create_cell_mappings, gen_default_cells, create_default_cell
)
from sdp_lib.passport.text_messages import Text
from sdp_lib.utils_common.utils_common import  timed


entity_patterns_and_aliases = (
    (PatternsDirectionTable.vehicle.value, DirectionEntities.vehicle,),
    (PatternsDirectionTable.arrow.value, DirectionEntities.arrow,),
    (PatternsDirectionTable.pedestrian.value, DirectionEntities.pedestrian,),
    (PatternsDirectionTable.always_red.value, DirectionEntities.always_red,),
    (PatternsDirectionTable.public.value, DirectionEntities.public,),
    (PatternsDirectionTable.tram.value, DirectionEntities.tram,),
)


def validate_tlc(direction_data: DirectionDataContainer, cell_mapping: CellMapping):
    ms = MessageStorage([], [])
    src_txt = cell_mapping.cell.text # Пример: "Тр. 7,8,9,10"
    repaired_text = repair_string_if_sep_in_illegal_pos(src_txt, ',')
    tlc_entity_is_valid = tlc_nums_is_valid = True
    if len(repaired_text) != len(src_txt):
        ms.add_errors(Text.illegal_pos_for_char(','))
    if (pos_start_nums_tlc:= found_pos_start_num(repaired_text)) is None:
        tlc_entity, nums = repaired_text, ''
        ms.add_errors(Text.has_not_num_tlc)
        tlc_nums_is_valid = False
    else:
        tlc_entity = repaired_text[:pos_start_nums_tlc] # Из примера: "Тр. "
        nums = repaired_text[pos_start_nums_tlc:].replace(' ', '') # Из примера: "7,8,9,10"
    if bad_nums := [n for n in nums.split(',') if not n.isdigit()]:
        tlc_nums_is_valid = False
        ms.add_errors(Text.invalid_nums(bad_nums))
    if direction_data.entity is not None:
        if not any(re.search(p, tlc_entity) for p in direction_data.tl_patterns):
            tlc_entity_is_valid = False
            ms.add_errors(Text.invalid_type_tlc)
    else:
        tlc_entity_is_valid = False
    res = bool(tlc_entity_is_valid and tlc_nums_is_valid)
    return CellData(
        value=src_txt,
        text_is_valid=res,
        context_is_valid=res,
        recovered_val=repaired_text,
        messages=ms,
        cell_mapping=cell_mapping
    )


def validate_timings_and_create_cell_data(cell_mapping: CellMapping, direction_data: DirectionDataContainer, t_name: str):
    ms = MessageStorage([], [])
    txt = cell_mapping.cell.text
    text_is_valid = False
    try:
        val_i = int(txt)
        text_is_valid = True
        if direction_data.entity is None:
            return CellData(
                txt, text_is_valid, None, converted_val=val_i, messages=ms, cell_mapping=cell_mapping
            )
    except ValueError:
        ms.add_errors(Text.is_not_a_number)
        return CellData(txt, text_is_valid, text_is_valid, messages=ms, cell_mapping=cell_mapping)
    # values: AllowedValues = timing_matches[(direction_data.entity, t_name)]
    values: AllowedValues = timing_matches[(direction_data.entity, t_name)]

    if values.min <= val_i <= values.max:  # OK case
        return CellData(txt, text_is_valid, text_is_valid, converted_val=val_i, messages=ms, cell_mapping=cell_mapping)

    if val_i < values.min:
        err = Text.val_must_be_gt(values.min)
    elif val_i > values.max:
        err = Text.val_must_be_lt(values.max)
    else:
        raise Exception(f'Debug: val_to_validate not fully validated')
    ms.add_errors(err)
    return CellData(txt, text_is_valid, False, messages=ms, cell_mapping=cell_mapping)


def validate_always_red_col_and_create_cell_data(
    cell_mapping: CellMapping,
    direction_entity: DirectionEntities,
):
    if direction_entity is None:
        return create_default_cell(cell_mapping)
    ms = MessageStorage([], [])
    src_txt = cell_mapping.cell.text
    if ((direction_entity != DirectionEntities.always_red and src_txt == '')
        or
        (direction_entity == DirectionEntities.always_red and re.match(PatternsDirectionTable.always_red_text_yes.value, src_txt))
    ): # ОК
        return CellData(
            value=src_txt,
            text_is_valid=True,
            context_is_valid=True,
            messages=ms,
            cell_mapping=cell_mapping
        )
    red_pattern = re.match(PatternsDirectionTable.always_red_text_yes.value, src_txt)
    if ((direction_entity != DirectionEntities.always_red and red_pattern)
        or
        (src_txt and red_pattern is None)
    ):
        ms.add_errors(Text.invalid_value)
    return CellData(
        value=src_txt,
        text_is_valid=bool(src_txt == '' or red_pattern is not None),
        context_is_valid=bool(ms.errors),
        messages=ms,
        cell_mapping=cell_mapping
    )


def validate_toov(cell_mapping: CellMapping, direction_data: DirectionDataContainer):
    if direction_data.entity is None:
        return create_default_cell(cell_mapping)
    txt = cell_mapping.cell.text
    text_is_valid = any(re.match(p, txt) is not None for p in direction_data.toov_patterns)

    return CellData(
        value=txt,
        text_is_valid=text_is_valid,
        context_is_valid=False if not text_is_valid else None,
        cell_mapping=cell_mapping,
        messages=MessageStorage([Text.invalid_value_for_direction_entity] if not text_is_valid else [], [])
    )


class DirectionTablePositionMapping(IntEnum):
    num                 = 0
    entity              = 1
    stages              = 2
    tlc                 = 3
    t_green_ext         = 4
    t_green_flashing    = 5
    t_green_yellow      = 6
    t_red               = 7
    t_red_yellow        = 8
    t_z                 = 9
    t_zz                = 10
    always_red          = 11
    toov_red            = 12
    toov_green          = 13
    description         = 14


# @timed
# def validate_directions_table(i_table: int, table: Table, ):
#     rows: _Rows = table.rows
#     geometry_check_list = validate_geometry(rows, allowed_column_lengths_dt, allowed_min_num_rows)
#     length = len(table.columns)
#     print(f'length: {length}')
#     names_and_patterns: HeadRowsDirectionTableData = dt_mapping_from_length[length]
#     print(names_and_patterns.second_row_names)
#     first_row = tuple(create_default_cells(i_table, 0, rows[0].cells))
#     second_row = tuple(create_cells_for_head_row(
#         i_table, 1, rows[1].cells, names_and_patterns.second_row_patterns, names_and_patterns.second_row_names
#     ))
#
#     print(f'length: {length}')
#     timing_columns = dt_timing_columns_mapping[length]
#     # print(first_row)
#     for i in range(2, len(rows)):
#         print(f'i: {i}')
#         gen_cell_mappings = gen_cell_mappings_and_lrstrip_in_cell_text(i_table, i, rows[i].cells)
#         num = num_validate_and_create_cell(
#             next(gen_cell_mappings)
#         ).write_messages_to_table_cell()
#         entity = match_cells_one_string_to_many_patterns_and_create_cell(
#             next(gen_cell_mappings),
#             entity_patterns_and_aliases,
#             True,
#         ).write_messages_to_table_cell()
#         stages = validate_sequence_directions_or_stages_nums_and_create_cell(
#             next(gen_cell_mappings)
#         ).write_messages_to_table_cell()
#
#         tlc = validate_tlc(entity.converted_val, next(gen_cell_mappings)).write_messages_to_table_cell()
#
#         # timings = (validate_timings(entity.recovered_val, col_name, next(cells)) if entity.recovered_val == DirectionEntities.vehicle else CellData('PLUG') for col_name in timing_columns)
#         timings = (
#             validate_timings(next(gen_cell_mappings), entity.converted_val, col_name,).write_messages_to_table_cell()
#             for col_name in timing_columns
#         )
#
#         # tzd = validate_number_and_create_cell((entity.recovered, ColNamesDirectionsTable.t_green_ext), next(cells))
#         # res = (num, entity, stages, tlc, tzd) + tuple(CellData('PLUG') for _ in range(9))
#         chain = itertools.chain(
#             (num, entity, stages, tlc),
#             timings,
#            (CellData('PLUG') for _ in range(4)),
#
#         )
#         r =  DirectionRow(tuple(c for c in chain))
#         print(r.represent(attr_splitter='\n') if i in (8, length - 100) else r)
#
#     doc.save('cadabra.docx')


# class _Validation:
#     def __init__(self, i_table: int, table: Table, ):
#         self.direction_rows: Sequence[DirectionRow] = []
#         self.i_table = i_table
#         self.table = table
#         self.rows: _Rows = table.rows
#         self.geometry_check_list = validate_geometry(self.rows, allowed_column_lengths_dt, allowed_min_num_rows)
#         self.length = len(table.columns)
#         self.first_row = lrstrip_in_cell_and_create_cell_mappings(i_table, 0, self.rows[0].cells)
#         names_and_patterns_second_row: HeadRowsDirectionTableData = dt_mapping_from_length[self.length]
#         self.second_row = tuple(create_cells_for_head_row(
#             self.i_table,
#             1,
#             self.rows[1].cells,
#             names_and_patterns_second_row.second_row_patterns,
#             names_and_patterns_second_row.second_row_names
#         ))
#         self.timing_columns = dt_timing_columns_mapping[self.length]
#         self.i_alw_red, self.i_toov_red, self.i_toov_green, self.i_description = range(self.length - 4, self.length)
#

class Structure(NamedTuple):
    num: int
    direction_entity: int
    stages: int
    tc: int
    t_green_extension: int
    t_green_flashing: int
    t_yellow: int
    t_red: int
    t_red_yellow: int
    t_z: int
    t_zz: int | None
    can_be_always_red: int
    toov_red: int
    toov_green: int
    description: int




@timed
def validate_directions_table(i_table: int, table: Table, ):
    rows: _Rows = table.rows
    # geometry_check_list = validate_geometry(table, allowed_column_lengths_dt, allowed_min_num_rows)


    the_table = TheTable(
        i_table,
        table,
        validate_geometry(table, allowed_column_lengths_dt, allowed_min_num_rows)
    )

    if not the_table.geometry_check_list.is_valid:
        for err in the_table.geometry_check_list.get_errors():
            the_table.messages.add_errors(err)
        target = table.add_row()
        target.cells[0].text = '\n'.join(f"*{e}" for e in the_table.messages.chain())
        target.cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 0, 0)
        doc.save('cadabra.docx')
        return table

        # names_and_patterns: HeadRowsDirectionTableData = head_rows_data_dt
    # first_row = lrstrip_in_cell_and_create_cell_mappings(i_table, 0, rows[0].cells)
    # first_row = tuple(gen_default_cells(first_row[0]))
    second_row = tuple(create_cells_for_head_row(
        i_table, 1, rows[1].cells, head_rows_data_dt.row1_col_patterns, head_rows_data_dt.row1_col_names
    ))
    """ 
    В данном блоке можно реализовать логику обработки валидности названий ячеек в колонках
    шапки таблицы(row[0] и row[1]).
    Например, прерывать валидацию, если присутствуют некорректные названия.
    if not all(c.context_is_valid for c in second_row): return
    """

    if not all(c.context_is_valid for c in second_row):
        doc.save('cadabra.docx')
        return
    return
    timing_columns = dt_timing_columns_mapping[length]
    i_alw_red, i_toov_red, i_toov_green, i_description = range(length - 4, length)
    for i in range(2, len(rows)):
        print(f'i: {i}')
        # gen_cell_mappings = gen_cell_mappings_and_lrstrip_in_cell_text(i_table, i, rows[i].cells)
        cell_mappings, empty_cells = lrstrip_in_cell_and_create_cell_mappings(i_table, i, rows[i].cells)
        if len(cell_mappings) == empty_cells:
            r = DirectionDataRow(tuple(gen_default_cells(c) for c in cell_mappings))
        else:
            num = num_validate_and_create_cell(cell_mappings[0]).write_messages_to_table_cell()
            entity = match_one_string_to_many_patterns_and_get_alias_and_create_cell(
                cell_mappings[1],
                entity_patterns_and_aliases,
                True,
            ).write_messages_to_table_cell()
            direction_data: DirectionDataContainer = mapping_direction_data.get(entity.converted_val)
            stages = validate_sequence_directions_or_stages_nums_and_create_cell(
                cell_mappings[2]
            ).write_messages_to_table_cell()
            tl = validate_tlc(direction_data, cell_mappings[3]).write_messages_to_table_cell()
            timings: Generator[CellData, Any, None] = (
                validate_timings_and_create_cell_data(
                    cell_mappings[ii],
                    entity.converted_val,
                    col_name,
                ).write_messages_to_table_cell()
                for ii, col_name in enumerate(timing_columns, 4)
            )
            always_red = validate_always_red_col_and_create_cell_data(
                cell_mappings[i_alw_red], entity.converted_val
            ).write_messages_to_table_cell()
            toov_red = validate_toov(
                cell_mappings[i_toov_red], entity.converted_val
            ).write_messages_to_table_cell()
            toov_green = validate_toov(
                cell_mappings[i_toov_green], entity.converted_val
            ).write_messages_to_table_cell()
            description = create_default_cell(cell_mappings[i_description])
            chain = itertools.chain(
                (num, entity, stages, tl),
                timings,
                (always_red, toov_red, toov_green, description),
                # (CellData('PLUG') for _ in range(3)),
            )
            r = DirectionDataRow(tuple(c for c in chain))
            print(r.represent(attr_splitter='\n') if i in (11, length - 100) else r)
        # print(r.represent(attr_splitter='\n') if i in (8, length - 100) else r)

    doc.save('cadabra.docx')


if __name__ == '__main__':
    # strings = ('1,2,2,4', '1.1,1.4,5,7,10', '', '     ', '1e,2dqd')
    # for s in strings:
    #     rr = check_directions_or_stages_string(s)
    #
    # # path = 'C://Programms//py.projects//sdp_lib//sdp_lib//passport//СО_2094_ул_Островитянова_ул_Ак_Волгина (2)'
    path1 = '/home/auser/Downloads/СО_2120_Северный_б_р_Санникова_ул_Декабристов_ул_'
    path2 = '/home/auser/Downloads/ПД Паспорт шаблон 2025 (Копия)'
    path3 = '/home/auser/Downloads/СО_2120_Северный_б_р_Санникова_ул_Декабристов_ул_ (1)'
    path4 = "C:\Programms\py.projects\sdp_lib\sdp_lib\passport\СО_2094_ул_Островитянова_ул_Ак_Волгина (2).docx"
    path5 = '/home/auser/py.projects/sdp_lib/sdp_lib/passport/СО_2094_ул_Островитянова_ул_Ак_Волгина_2.docx'
    path6 = '/home/auser/py.projects/sdp_lib/sdp_lib/passport/passport2/validation/ПД Паспорт шаблон 2025.docx'


    doc = Document(path5)
    # c = CheckListTable()
    validate_directions_table(0, doc.tables[0])

