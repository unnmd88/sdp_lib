import asyncio
import time
from collections.abc import Sequence
from typing import NamedTuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pysnmp.smi.rfc1902 import ObjectType, ObjectIdentity

from sdp_lib.management_controllers.snmp import snmp_utils
from sdp_lib.management_controllers.snmp.oids import Oids
from sdp_lib.management_controllers.snmp.snmp_api import PotokP
from sdp_lib.management_controllers.snmp.snmp_requests import snmp_engine

from openpyxl import Workbook

from sdp_lib.utils_common import utils_common
from sdp_lib.utils_common.utils_common import get_curr_datetime


class TactValue(NamedTuple):
    as_int: int
    as_hex: str
    duration: float


def write_to_excel(name_wb, wb, sheet, cycles):

    row = sheet.max_row + 2
    for i, tact_data in enumerate(cycles[-1], 1):
        cell = sheet.cell(row=row, column=i)
        cell.value = tact_data.as_int
        cell.alignment = Alignment(horizontal='center')
        if tact_data.as_int == 2:
            cell.fill = PatternFill(fill_type='solid', fgColor='00FF00')
        elif tact_data.as_int == 1:
            cell.fill = PatternFill(fill_type='solid', fgColor='FF7F00')
        cell = sheet.cell(row=row + 1, column=i)
        cell.value = tact_data.as_hex
        cell.alignment = Alignment(horizontal='center')

        cell = sheet.cell(row=row + 2, column=i)
        cell.value = tact_data.duration
        cell.alignment = Alignment(horizontal='center')

    wb.save(name_wb)


# async def main(timeout: float = 0.1):
#     wb = Workbook()
#     ws = wb.active
#     host = PotokP(ipv4='10.45.154.12', engine=snmp_engine)
#     await host.get_scn_from_host_and_set_to_attr()
#     # _oids = [ObjectType(ObjectIdentity(oid)) for oid in [Oids.utcReplyGn, Oids.potokP_utcReplyPromTact]]
#     _oids = [
#         Oids.utcReplyGn + host.scn.scn_as_ascii,
#         Oids.potokP_utcReplyPromTact + host.scn.scn_as_ascii
#     ]
#
#     print(f'SCN WAS SET: {host.scn.scn_as_ascii}')
#
#     res = await host.request_sender.snmp_get([ObjectType(ObjectIdentity(oid)) for oid in _oids])
#
#     prev_stage_as_hex = res[3][0][1].prettyPrint()
#     prev_stage_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_stage_as_hex)
#
#     prev_prom_as_hex = res[3][1][1].prettyPrint()
#     prev_prom_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_prom_as_hex)
#
#     tmr1 = time.perf_counter()
#     cycles = []
#     prom_sequence = []
#
#     while True:
#         err_indication, err_status, err_index, varbinds = await host.request_sender.snmp_get(
#             [ObjectType(ObjectIdentity(oid)) for oid in _oids]
#         )
#         # stg = f'{res[3][0][0]}(Stage)={res[3][0][1].prettyPrint()}({snmp_utils.convert_val_as_hex_to_decimal(res[3][0][1].prettyPrint())})'
#         # prom = f'{res[3][1][0]}(Prom)={res[3][1][1].prettyPrint()}({snmp_utils.convert_val_as_hex_to_decimal(res[3][1][1].prettyPrint())})'
#         if varbinds:
#             curr_stage_as_hex = varbinds[0][1].prettyPrint()
#             curr_prom_as_hex = varbinds[1][1].prettyPrint()
#
#             curr_stage_as_int = snmp_utils.convert_val_as_hex_to_decimal(curr_stage_as_hex)
#             curr_prom_as_int = snmp_utils.convert_val_as_hex_to_decimal(curr_prom_as_hex)
#
#             # print(f'{curr_prom_as_hex=} {curr_prom_as_int=}')
#             if prev_prom_as_hex != curr_prom_as_hex:
#                 duration = time.perf_counter() - tmr1
#                 tmr1 = time.perf_counter()
#                 # prom_sequence += (prev_prom_as_int, prev_prom_as_hex, duration)
#                 prom_sequence.append((f'{prev_prom_as_int}({prev_prom_as_hex})', f'{duration}.c'))
#                 print(f'stg={prev_stage_as_hex}({prev_stage_as_int}) prom={prev_prom_as_hex}({prev_prom_as_int}) | duration={duration}')
#                 if curr_stage_as_int == 1 and curr_prom_as_int == 1:
#                     cycles.append(prom_sequence)
#                     print(prom_sequence)
#                     asyncio.create_task(asyncio.to_thread(write_to_excel, wb, ws, cycles))
#                     prom_sequence = []
#
#             prev_stage_as_hex = curr_stage_as_hex
#             prev_stage_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_stage_as_hex)
#             prev_prom_as_hex = curr_prom_as_hex
#             prev_prom_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_prom_as_hex)
#
#         await asyncio.sleep(timeout)


async def main(timeout: float = 0.1):
    wb_name = 'report_prom_oid_ug405.xlsx'
    try:
        wb = openpyxl.load_workbook(wb_name)
    except FileNotFoundError:
        wb = Workbook()
        wb.save(wb_name)



    # host = PotokP(ipv4='10.45.154.12', engine=snmp_engine)
    host = PotokP(ipv4='10.179.65.153', engine=snmp_engine)
    await host.get_scn_from_host_and_set_to_attr()
    try:
        ws = wb[f'{host.ip_v4}_{host.scn.scn_as_chars}']
    except KeyError:
        ws = wb.create_sheet(f'{host.ip_v4}_{host.scn.scn_as_chars}', 0)


    cell = ws.cell(row=ws.max_row + 2, column=1)
    cell.value = f'Время начала отчёта: {get_curr_datetime()}'
    # _oids = [ObjectType(ObjectIdentity(oid)) for oid in [Oids.utcReplyGn, Oids.potokP_utcReplyPromTact]]
    _oids = [
        Oids.utcReplyGn + host.scn.scn_as_ascii,
        Oids.potokP_utcReplyPromTact + host.scn.scn_as_ascii
    ]

    print(f'SCN WAS SET: {host.scn.scn_as_ascii}')

    res = await host.request_sender.snmp_get([ObjectType(ObjectIdentity(oid)) for oid in _oids])

    prev_stage_as_hex = res[3][0][1].prettyPrint()
    prev_stage_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_stage_as_hex)

    prev_prom_as_hex = res[3][1][1].prettyPrint()
    prev_prom_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_prom_as_hex)

    tmr1 = time.perf_counter()
    cycles = []
    prom_sequence = []

    while True:
        err_indication, err_status, err_index, varbinds = await host.request_sender.snmp_get(
            [ObjectType(ObjectIdentity(oid)) for oid in _oids]
        )
        # stg = f'{res[3][0][0]}(Stage)={res[3][0][1].prettyPrint()}({snmp_utils.convert_val_as_hex_to_decimal(res[3][0][1].prettyPrint())})'
        # prom = f'{res[3][1][0]}(Prom)={res[3][1][1].prettyPrint()}({snmp_utils.convert_val_as_hex_to_decimal(res[3][1][1].prettyPrint())})'
        if varbinds:
            curr_stage_as_hex = varbinds[0][1].prettyPrint()
            curr_prom_as_hex = varbinds[1][1].prettyPrint()

            curr_stage_as_int = snmp_utils.convert_val_as_hex_to_decimal(curr_stage_as_hex)
            curr_prom_as_int = snmp_utils.convert_val_as_hex_to_decimal(curr_prom_as_hex)

            # print(f'{curr_prom_as_hex=} {curr_prom_as_int=}')
            if prev_prom_as_hex != curr_prom_as_hex:
                duration = time.perf_counter() - tmr1
                tmr1 = time.perf_counter()
                # prom_sequence += (prev_prom_as_int, prev_prom_as_hex, duration)
                # prom_sequence.append((f'{prev_prom_as_int}({prev_prom_as_hex})', f'{duration}.c'))
                prom_sequence.append((TactValue(prev_prom_as_int, prev_prom_as_hex, round(duration, 3))))
                print(f'stg={prev_stage_as_hex}({prev_stage_as_int}) prom={prev_prom_as_hex}({prev_prom_as_int}) | duration={duration}')
                if curr_stage_as_int == 1 and curr_prom_as_int == 1:
                    cycles.append(prom_sequence)
                    print(prom_sequence)
                    asyncio.create_task(asyncio.to_thread(write_to_excel, wb_name, wb, ws, cycles))
                    prom_sequence = []

            prev_stage_as_hex = curr_stage_as_hex
            prev_stage_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_stage_as_hex)
            prev_prom_as_hex = curr_prom_as_hex
            prev_prom_as_int = snmp_utils.convert_val_as_hex_to_decimal(prev_prom_as_hex)

        await asyncio.sleep(timeout)


if __name__ == '__main__':

    asyncio.run(main())
