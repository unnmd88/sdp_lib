from collections.abc import Iterable

import openpyxl
from openpyxl import Workbook


DET_HEAD = ['Начало', 'Конец', "Состояние", 'Детектор', 'Продолжительность']


class ExcelLogger:

    def __init__(
            self,
            wb_name: str,
            sheet_name: str,
            ip: str = '',
            num_co: str = '',
            head: Iterable[str] = None,
    ):
        self.ip = ip
        self.num_co = num_co
        _name = wb_name.split('.')
        if len(_name[-1]) == 0:
            raise ValueError(f'wb_name cant be empty')
        elif _name[-1].endswith('xlsx'):
            self.wb_name = wb_name
        else:
            self.wb_name = f'{wb_name}.xlsx'

        self.sheet_name = sheet_name
        try:
            self._wb = openpyxl.load_workbook(wb_name)
        except FileNotFoundError:
            self._wb = Workbook()

        try:
            self.ws = self._wb[self.sheet_name]
        except KeyError:
            for sheet_name in self._wb.sheetnames:
                sheet = self._wb.get_sheet_by_name(sheet_name)
                self._wb.remove_sheet(sheet)
            self.ws = self._wb.create_sheet(self.sheet_name, 0)
        if not self.ws['A1'].value and head is not None:
            for col, data in enumerate(head, 1):
                self.ws.cell(row=1, column=col).value = data

        self._wb.save(wb_name)

    def append_to_next_row(self, data: Iterable):
        row = self.ws.max_row + 1
        for col, data in enumerate(data, 1):
            self.ws.cell(row=row, column=col).value = data
        self._wb.save(self.wb_name)

