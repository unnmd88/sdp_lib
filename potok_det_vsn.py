import asyncio

import openpyxl
from openpyxl import Workbook
from pysnmp.smi.rfc1902 import ObjectType, ObjectIdentity

from excel_logs import ExcelLogger, DET_HEAD
from sdp_lib.management_controllers.snmp.oids import Oids
from sdp_lib.management_controllers.snmp.snmp_core import PotokP
from sdp_lib.management_controllers.snmp.snmp_requests import snmp_engine
from sdp_lib.management_controllers.snmp.snmp_utils import PrettyVsnString
from observers import BaseSwitchStateObserver
import logging

logger = logging.getLogger(__name__)
logger.setLevel(logging.DEBUG)
formatter = logging.Formatter('{name} [{levelname:^11}] {asctime} {message}', style='{')
file_handler = logging.FileHandler('vsn.log')
file_handler.setFormatter(formatter)
logger.addHandler(file_handler)


def create_xlsx(host: PotokP):
    wb_name = 'utcReplyVSn_log.xlsx'
    try:
        wb = openpyxl.load_workbook(wb_name)
    except FileNotFoundError:
        wb = Workbook()
        wb.save(wb_name)
    try:
        ws = wb[f'{host.ip_v4}_{host.scn.scn_as_chars}']
    except KeyError:
        for sheet_name in wb.sheetnames:
            sheet = wb.get_sheet_by_name(sheet_name)
            wb.remove_sheet(sheet)
        ws = wb.create_sheet(f'{host.ip_v4}_{host.scn.scn_as_chars}', 0)
    if ws.max_row < 1:
        ws.append([host.ip_v4, host.scn.scn_as_chars])
    return wb


def create_sheet(host: PotokP, wb: Workbook):
    try:
        wb[f'{host.ip_v4}_{host.scn.scn_as_chars}']
    except KeyError:
        wb.create_sheet(f'{host.ip_v4}_{host.scn.scn_as_chars}', 0)


def write_to_excel(name_wb, wb, sheet, cycles):

    row = sheet.max_row + 2
    for i, tact_data in enumerate(cycles[-1], 1):
        cell = sheet.cell(row=row, column=i)
        cell.value = tact_data.as_int
        cell.alignment = Alignment(horizontal='center')
        if tact_data.as_int == 2:
            cell.fill = PatternFill(fill_type='solid', fgColor='00FF00')
        elif tact_data.as_int == 1:
            cell.fill = PatternFill(fill_type='solid', fgColor='FF7F00')
        cell = sheet.cell(row=row + 1, column=i)
        cell.value = tact_data.as_hex
        cell.alignment = Alignment(horizontal='center')

        cell = sheet.cell(row=row + 2, column=i)
        cell.value = tact_data.duration
        cell.alignment = Alignment(horizontal='center')

    wb.save(name_wb)


async def main(timeout: float = .8):

    ipv4 = '10.179.72.97'
    snmp_host = PotokP(ipv4=ipv4, engine=snmp_engine, host_id='190')

    await snmp_host.get_scn_from_host_and_set_to_attr()
    print(f'host ip: {snmp_host.ip_v4}')
    print(f'host id: {snmp_host.host_id}')
    print(f'scn from host as chars: {snmp_host.scn.scn_as_chars}')
    print(f'scn from host as ASCII: {snmp_host.scn.scn_as_ascii}')

    excel_logger = ExcelLogger(
        wb_name='utcReplyVSn_log.xlsx',
        sheet_name=f'{snmp_host.ip_v4}_{snmp_host.scn.scn_as_chars}',
        ip=snmp_host.ip_v4,
        num_co=snmp_host.scn.scn_as_chars,
        head=DET_HEAD
    )

    samples = await snmp_host.request_sender.snmp_get([ObjectType(ObjectIdentity('.1.3.6.1.4.1.13267.3.2.2.5.0'))])
    print(f'Samples: {samples[3][0][1].prettyPrint()}')
    pretty_vsn = PrettyVsnString(int(samples[3][0][1].prettyPrint()))
    det_cnt = await snmp_host.request_sender.snmp_get([ObjectType(ObjectIdentity(Oids.utcType2ScootDetectorCount + snmp_host.scn.scn_as_ascii))])
    print(f'NumDetectors: {det_cnt[3][0][1].prettyPrint()}')

    _oids = [ObjectType(ObjectIdentity(Oids.utcReplyVSn + snmp_host.scn.scn_as_ascii))]

    observer1 = BaseSwitchStateObserver(name='D1', excel_logger=excel_logger)
    observer2 = BaseSwitchStateObserver(name='D2', excel_logger=excel_logger)
    try:
        while True:
            err_indication, err_status, err_index, varbinds = await snmp_host.request_sender.snmp_get(
                _oids,
                timeout=.6,
            )
            if err_indication or not varbinds:
                observer1.check(None)
                observer2.check(None)
            else:
                octet_string_vsn = varbinds[0][1].prettyPrint()
                print(f'PrettyOctStr: {pretty_vsn.get_pretty_string(octet_string_vsn)}')
                observer1.check_and_write_log(octet_string_vsn[3])
                observer2.check_and_write_log(octet_string_vsn[2])
            await asyncio.sleep(timeout)
    except Exception as e:
        logger.critical('FAIL!!!!')









if __name__ == '__main__':
    asyncio.run(main())